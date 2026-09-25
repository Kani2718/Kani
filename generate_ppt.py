import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

# Create a presentation
prs = Presentation()

# Helper to add a title slide
def add_title_slide(title, subtitle=""):
    slide_layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    if subtitle:
        slide.placeholders[1].text = subtitle

# Helper to add bullet slide
def add_bullet_slide(title, bullet_points):
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
        p.font.size = Pt(18)

# Slides
add_title_slide("Smart Campus AI", "Premium Authentication & User Flow")

add_bullet_slide("Overview", [
    "Goal: Modern, production‑quality SaaS authentication experience.",
    "Premium UI with split‑screen layout and glassmorphic cards.",
    "Full client‑side & server‑side validation.",
    "Secure session management & role‑based access."
])

add_bullet_slide("Authentication Flow", [
    "New User: Landing → Register → Success → Auto‑redirect to Login → Login → Dashboard.",
    "Existing User: Landing → Login → Dashboard.",
    "After registration: show success banner, pre‑fill email, do NOT auto‑open dashboard."
])

add_bullet_slide("Premium Login UI", [
    "Split‑screen: left visual, right auth card.",
    "Glassmorphic card with backdrop blur, rounded corners.",
    "Fields: Email, Password, Remember Me, Forgot Password.",
    "Demo Account button for hackathon."
])

add_bullet_slide("Premium Register UI", [
    "Same split‑screen layout.",
    "Fields: Full Name, Email, Phone, Institution, Campus, Role, Password, Confirm Password.",
    "Terms & Privacy checkbox.",
    "Password strength meter (Weak → Medium → Strong).",
    "Success screen with animation and auto‑redirect countdown."
])

add_bullet_slide("Validation Rules", [
    "Full Name: required, ≥2 chars.",
    "Email: required, valid format, unique.",
    "Phone: required, valid format.",
    "Password: ≥8 chars, upper, lower, number, special.",
    "Confirm Password must match.",
    "Role selection required.",
    "Real‑time client validation + server checks."
])

add_bullet_slide("Routing & Protection", [
    "Public routes: /, /about, /features, /login, /register, /forgot-password.",
    "Protected routes: /dashboard, /classrooms, /energy, ….",
    "Unauthenticated access → redirect to /login.",
    "After login, redirect back to originally requested page."
])

add_bullet_slide("Session Management", [
    "Secure password hashing.",
    "JWT or secure cookies for auth tokens.",
    "Session stored in localStorage/sessionStorage.",
    "Logout clears session and redirects to login.",
    "Role‑based authorization for protected pages."
])

add_bullet_slide("Demo Account (Hackathon)", [
    "Email: demo@smartcampus.ai",
    "Password: configurable via env vars.",
    "Button fills login form and authenticates normally."
])

add_bullet_slide("Logout Flow", [
    "Profile dropdown includes Logout option.",
    "Confirmation modal with Cancel / Sign Out.",
    "On sign out: destroy session, show success toast, redirect to login."
])

add_bullet_slide("Branding & Micro‑Interactions", [
    "Deep navy/black gradient background with subtle glows.",
    "Glassmorphic auth cards, modern typography (Inter/Manrope).",
    "Animations: input focus glow, button hover/press, password visibility toggle, success checkmark pop, error shake, card entrance."
])

add_bullet_slide("Responsive Design", [
    "Desktop: 50/50 split‑screen.",
    "Tablet & mobile: hide illustration, full‑width form, large touch targets, no horizontal scroll.",
    "Consistent branding across all auth pages."
])

# Save PPT
output_path = Path(__file__).parent / "Smart_Campus_AI_Auth_Flow.pptx"
prs.save(str(output_path))
print(f"Presentation saved to {output_path}")
